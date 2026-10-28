"""Regression tests for the Reports/Overview fixes: true 7-day trends,
like-for-like comparison windows, neutral 0.0% arrows, cent-exact money,
inclusive report titles, accent-folded stock search, numeric Excel cells
and Turkish-capable PDFs."""

import shutil
import subprocess
import tempfile
from datetime import date, datetime
from pathlib import Path
from unittest import mock

from shared import analytics
from shared.builders.report_builder import (
    NumericText, ReportDocument, ReportSection, build_sales_report, line_cents,
)
from shared.models import Dealership, LineItem, Product, StockLevel, StockLocation, Transaction

DEALERS = [Dealership("MET-01", "Metro Heavy", "Metro", "Columbus")]


def _sale(day: date, amount, code="MET-01", qty=1):
    return Transaction(
        items=[LineItem(product_barcode="B1", product_name_at_sale="Box", unit_price_at_sale=amount, quantity=qty)],
        created_at=datetime(day.year, day.month, day.day, 12), dealership_code=code,
    )


def _daily(start: date, end: date, amount):
    out, d = [], start
    while d <= end:
        out.append(_sale(d, amount))
        d = date.fromordinal(d.toordinal() + 1)
    return out


# --- 1: the 7-day week reaches back before the period ---------------------


def test_week_start_reaches_before_a_young_period():
    p = analytics.period_for("month", date(2026, 10, 5))
    assert p.week_start == date(2026, 9, 29)
    assert analytics.period_for("month", date(2026, 10, 20)).week_start == date(2026, 10, 1)


def test_steady_sales_across_a_month_boundary_trend_flat():
    today = date(2026, 10, 5)
    p = analytics.period_for("month", today)
    everything = _daily(date(2026, 9, 29), today, 100)
    in_period = [t for t in everything if t.created_at.date() >= p.start]
    top = analytics.top_dealerships(in_period, DEALERS, today, week_transactions=everything)
    assert top[0].week == [100.0] * 7 and top[0].trend == 0.0
    assert top[0].revenue == 500.0  # the ranking is still the period only
    view = analytics.build_report_view(p, in_period, [], DEALERS, week_transactions=everything)
    assert view.top[0].week == [100.0] * 7 and view.top[0].trend == 0.0
    report = analytics.build_period_report(in_period, p, DEALERS, week_transactions=everything)
    top_rows = [s for s in report.sections if s.title.startswith("Top dealerships")][0].rows
    assert top_rows[0][4] == "0.0%"


# --- 3: like-for-like comparison windows ----------------------------------


def test_31_march_compares_28_days_with_february():
    p = analytics.period_for("month", date(2026, 3, 31))
    assert (p.prev_start, p.prev_end) == (date(2026, 2, 1), date(2026, 2, 28))
    assert p.comparable_end == date(2026, 3, 28) and p.comparable_days == 28 and p.is_clamped
    assert p.days == 31  # the headline still covers the whole window


def test_unclamped_period_compares_everything():
    p = analytics.period_for("month", date(2026, 9, 24))
    assert p.compare_end is None and p.comparable_end == p.end and not p.is_clamped
    assert (p.prev_end - p.prev_start).days == (p.comparable_end - p.start).days


def test_equal_daily_sales_in_march_and_february_change_by_zero():
    p = analytics.period_for("month", date(2026, 3, 31))
    march = _daily(date(2026, 3, 1), date(2026, 3, 31), 100)
    feb = _daily(date(2026, 2, 1), date(2026, 2, 28), 100)
    view = analytics.build_report_view(p, march, feb, DEALERS)
    assert view.revenue == 3100.0 and view.previous_revenue == 2800.0 and view.change == 0.0
    assert view.projected_total is None  # 31 Mar closes the month
    totals = dict(analytics.build_period_report(march, p, DEALERS, previous=feb).sections[0].rows)
    assert totals["Change"] == "0.0%" and "Change compares" in totals


def test_leap_year_ytd_is_clamped_on_both_sides():
    p = analytics.period_for("ytd", date(2024, 12, 31))  # 366 days vs 2023's 365
    assert p.prev_end == date(2023, 12, 31) and p.comparable_end == date(2024, 12, 30)


# --- 7: neutral arrow ------------------------------------------------------


def test_a_change_that_rounds_to_zero_is_neutral():
    assert analytics.percent_change(99.96, 100) == 0.0
    assert str(analytics.percent_change(99.96, 100)) == "0.0"  # not "-0.0"
    for tiny in (-0.04, -0.0, 0.04, 0.0):
        assert analytics.change_text(tiny) == "• 0.0%"
        assert analytics.change_text(tiny, plain=True) == "0.0%"
        assert analytics.change_direction(tiny) == 0
    assert analytics.change_text(-0.06) == "▼ 0.1%" and analytics.change_direction(-0.06) == -1
    assert analytics.change_text(0.06) == "▲ 0.1%" and analytics.change_direction(None) is None


# --- 8: money in cents -----------------------------------------------------


def test_three_decimal_prices_round_half_up_per_line():
    item = LineItem("B1", "Box", 1.005, 1)
    assert line_cents(item) == 101  # binary float rounding would give 100
    assert line_cents(LineItem("B1", "Box", 0.335, 3)) == 101
    sales = [_sale(date(2026, 9, 1), 1.005, qty=1) for _ in range(3)]
    assert analytics.daily_totals(sales, date(2026, 9, 1), date(2026, 9, 1)) == [3.03]
    assert analytics.revenue_by_region(sales, DEALERS)["Metro"] == 3.03
    assert analytics.sales_today(sales, date(2026, 9, 1)) == (3, 3.03)
    doc = build_sales_report(sales, datetime(2026, 9, 1), datetime(2026, 9, 2))
    assert dict(doc.sections[0].rows)["Revenue"] == "3.03"
    assert doc.sections[1].rows[0][3] == "3.03"


# --- 6: report title shows the inclusive end ------------------------------


def test_sales_report_title_prints_the_last_day_not_the_exclusive_bound():
    doc = build_sales_report([], datetime(2026, 10, 1), datetime(2026, 10, 4))  # 1st..3rd inclusive
    assert doc.title == "Sales Report: 2026-10-01 to 2026-10-03"
    one_day = build_sales_report([], datetime(2026, 10, 1), datetime(2026, 10, 2))
    assert one_day.title == "Sales Report: 2026-10-01 to 2026-10-01"


# --- 4: stock filter folds accents ----------------------------------------


def test_stock_search_is_accent_and_case_insensitive_on_both_sides():
    from shared import overview

    products = [
        Product("P1", "Çay Şeker", 1, stock_quantity=5, critical_stock_level=1),
        Product("P2", "İpek Kumaş", 1, stock_quantity=5, critical_stock_level=1),
        Product("P3", "Hose", 1, stock_quantity=5, critical_stock_level=1),
    ]
    names = lambda q: [r.name for r in overview.stock_grid(products, [], [], query=q).rows]  # noqa: E731
    assert names("cay") == ["Çay Şeker"] and names("CAY seker") == ["Çay Şeker"]
    assert names("ipek") == ["İpek Kumaş"] and names("İPEK") == ["İpek Kumaş"]
    assert names("çay") == ["Çay Şeker"] and names("kumas") == ["İpek Kumaş"]
    assert names("p3") == ["Hose"] and names("zzz") == []


# --- 5: Excel numbers are numbers -----------------------------------------


def test_excel_numeric_cells_are_numbers_and_text_stays_safe_text():
    from openpyxl import load_workbook

    from admin_app.export.excel_report_exporter import export_to_excel

    report = analytics.build_period_report(
        [_sale(date(2026, 9, 1), 100.0), _sale(date(2026, 9, 2), 50.5)],
        analytics.period_for("month", date(2026, 9, 3)), DEALERS,
        previous=[_sale(date(2026, 8, 2), 100.0)],
    )
    report.sections.append(ReportSection("Plain", [("0012345", "=cmd()", "12", "3.50", "+8.4%", "-3.2%", "no prior data")]))
    report.sections.append(ReportSection("Named", [("N1", "1907", "0.50", "-2+3 Kalem")]))
    with mock.patch("admin_app.export.excel_report_exporter.store_name", lambda: "=Store"):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.xlsx"
            export_to_excel(report, path)
            sheet = load_workbook(path).active
    cells = {}
    for row in sheet.iter_rows():
        label = row[0].value
        if label is not None:
            cells.setdefault(label, row)
    assert sheet["A1"].value == "=Store" and sheet["A1"].data_type == "s"  # shown as typed, never evaluated
    rev = cells["Revenue"][1]
    assert rev.value == 150.5 and isinstance(rev.value, float) and rev.number_format == "#,##0.00"
    sales = cells["Sales"][1]
    assert sales.value == 2 and isinstance(sales.value, int)
    change = cells["Change"][1]
    assert change.value == 0.505 and "%" in change.number_format
    plain = cells["0012345"]
    assert plain[0].data_type == "s" and plain[0].value == "0012345"  # first column stays text
    assert plain[1].value == "=cmd()" and plain[1].data_type == "s"  # a string cell: nothing to defuse, nothing runs
    # plain strings stay text, however numeric they look; only the builders' NumericText becomes a number
    assert [(c.value, c.data_type) for c in plain[2:6]] == [("12", "s"), ("3.50", "s"), ("+8.4%", "s"), ("-3.2%", "s")]
    assert plain[6].value == "no prior data"


def test_numeric_text_still_prints_as_its_text():
    cell = NumericText("150.00", 150.0, "#,##0.00")
    assert cell == "150.00" and f"{cell}" == "150.00" and cell.number == 150.0


# --- 2: PDF carries Turkish letters ---------------------------------------

TURKISH = "Şişli Şubesi Çay İçin ğüşıö"


def _pdf(report, store=TURKISH):
    from admin_app.export.pdf_report_exporter import export_to_pdf

    with mock.patch("admin_app.export.pdf_report_exporter.store_name", lambda: store):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.pdf"
            export_to_pdf(report, path)
            data = path.read_bytes()
            text = ""
            if shutil.which("pdftotext"):
                text = subprocess.run(
                    ["pdftotext", "-enc", "UTF-8", "-layout", str(path), "-"], capture_output=True
                ).stdout.decode("utf-8", errors="replace")
    return data, text


def test_pdf_embeds_a_unicode_font_and_keeps_turkish_glyphs():
    report = ReportDocument("Başlık ĞÜŞİÖÇ", datetime(2026, 10, 4, 9, 0), [
        ReportSection("Bölüm", [("Çay Şeker", NumericText("5", 5), "10.00"), ("İpek ığ", "3", "7.00")]),
    ])
    data, text = _pdf(report)
    assert b"DejaVuSans" in data and b"DejaVuSans-Bold" in data
    if text:
        for fragment in (TURKISH, "Başlık ĞÜŞİÖÇ", "Çay Şeker", "İpek ığ"):
            assert fragment in text


def test_pdf_falls_back_to_helvetica_when_the_font_is_missing():
    report = ReportDocument("Title", datetime(2026, 10, 4, 9, 0), [ReportSection("S", [("a", "1")])])
    with mock.patch("admin_app.export.pdf_report_exporter.report_font_files", lambda: None):
        data, _ = _pdf(report, store="Shop")
    assert b"DejaVu" not in data and data.startswith(b"%PDF")


def test_pdf_long_names_are_truncated_not_run_off_the_page():
    from reportlab.pdfbase import pdfmetrics

    from admin_app.export import pdf_report_exporter as exporter

    regular, _ = exporter.register_fonts()
    fitted = exporter.fit_text("Ş" * 300, regular, 9, 100)
    assert fitted.endswith("…") and pdfmetrics.stringWidth(fitted, regular, 9) <= 100
    assert exporter.fit_text("short", regular, 9, 100) == "short"
    report = ReportDocument("T", datetime(2026, 10, 4), [ReportSection("S", [("X" * 400, "1", "2.00")] * 80)])
    data, text = _pdf(report)  # many long rows: paginates without error
    assert data.startswith(b"%PDF")
    if text:
        assert max(len(line) for line in text.splitlines()) < 200


def test_bundled_fonts_exist_and_ship_with_every_build():
    from shared.build_manifest import BUNDLED_DATA_FILES
    from shared.fonts import report_font_files

    files = report_font_files()
    assert files is not None and all(f.is_file() for f in files)
    shipped = {rel for rel, _ in BUNDLED_DATA_FILES}
    assert {"shared/assets/fonts/DejaVuSans.ttf", "shared/assets/fonts/DejaVuSans-Bold.ttf"} <= shipped
    assert (files[0].parent / "LICENSE-DejaVu.txt").is_file()


def test_excel_names_that_look_like_numbers_or_formulas_stay_plain_text(tmp_path):
    from openpyxl import load_workbook

    from admin_app.export.csv_report_exporter import export_to_csv
    from admin_app.export.excel_report_exporter import export_to_excel
    from shared.builders.report_builder import amount_text, count_text

    names = ("1907", "0.50", '=HYPERLINK("http://x","y")', "-2+3 Kalem", "Çay\x07 Şeker")
    report = ReportDocument("R", datetime(2026, 9, 3, 12, 0), [
        ReportSection("Per-Product Breakdown", [(f"B{i}", name, count_text(3), amount_text(12.5)) for i, name in enumerate(names)]),
    ])
    with mock.patch("admin_app.export.excel_report_exporter.store_name", lambda: "Shop"):
        export_to_excel(report, tmp_path / "r.xlsx")
    export_to_csv(report, tmp_path / "r.csv")

    rows = {r[0].value: r for r in load_workbook(tmp_path / "r.xlsx").active.iter_rows() if r[0].value}
    cells = [rows[f"B{i}"][1] for i in range(len(names))]
    assert [(c.value, c.data_type) for c in cells] == [
        ("1907", "s"), ("0.50", "s"), ('=HYPERLINK("http://x","y")', "s"), ("-2+3 Kalem", "s"), ("Çay Şeker", "s"),
    ]  # no leading apostrophe, control characters dropped
    units, revenue = rows["B0"][2], rows["B0"][3]
    assert (units.value, units.number_format) == (3, "#,##0") and units.data_type == "n"
    assert (revenue.value, revenue.number_format) == (12.5, "#,##0.00") and revenue.data_type == "n"
    assert "'=HYPERLINK" in (tmp_path / "r.csv").read_text(encoding="utf-8-sig")  # CSV keeps its defence

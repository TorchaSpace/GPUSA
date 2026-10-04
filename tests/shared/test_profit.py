"""Gross profit and margin in the analytics, report builders and exports:
only lines whose cost was known count, the unknown share is stated, money
is whole cents, and nothing divides by zero."""

import csv
import tempfile
from datetime import date, datetime
from pathlib import Path
from unittest import mock

from shared import analytics
from shared.builders.report_builder import (
    ProfitSummary, build_sales_report, margin_text, profit_summary,
)
from shared.models import Dealership, LineItem, Transaction

DEALERS = [
    Dealership("MET-01", "Metro Heavy", "Metro", "Columbus"),
    Dealership("CST-04", "Harbor Point", "Coastal", "Norfolk"),
]


def _line(barcode, price, qty, cost=None):
    """cost None = the sale predates costing (cost unknown)."""
    return LineItem(barcode, f"Name {barcode}", price, qty,
                    unit_cost_at_sale=cost or 0.0, cost_known=cost is not None)


def _sale(day, items, code="MET-01"):
    return Transaction(items=items, created_at=datetime(2026, 9, day, 12), dealership_code=code)


# A: 3 x 10.00 at cost 6.00   -> revenue 30.00, cost 18.00
# B: 2 x  5.50, cost unknown  -> revenue 11.00
# C: 1 x 20.00 at cost 12.50  -> revenue 20.00, cost 12.50
SALES = [
    _sale(1, [_line("A", 10.0, 3, 6.0), _line("B", 5.5, 2)]),
    _sale(2, [_line("C", 20.0, 1, 12.5)], code="CST-04"),
]


def test_profit_summary_numeric_example():
    p = profit_summary(SALES)
    assert (p.revenue_cents, p.known_revenue_cents, p.cost_cents, p.known_lines) == (6100, 5000, 3050, 2)
    assert p.profit_cents == 1950 and p.profit == 19.5
    assert p.margin == 39.0  # 19.50 / 50.00 - the unknown line's 11.00 is not in the denominator
    assert p.unknown_revenue_cents == 1100 and p.is_partial
    assert p.unknown_percent == 18  # 11 / 61 = 18.03 %


def test_cents_are_exact_where_floats_are_not():
    # 0.10 x 3 = 0.30000000000000004 and 0.07 x 3 = 0.21000000000000002 as floats
    p = profit_summary([_sale(1, [_line("A", 0.10, 3, 0.07)])])
    assert (p.revenue_cents, p.cost_cents, p.profit_cents) == (30, 21, 9)
    assert p.margin == 30.0
    # a 3-decimal cost rounds half-up per line: 1.005 x 1 -> 1.01
    q = profit_summary([_sale(1, [_line("A", 5.0, 1, 1.005)])])
    assert q.cost_cents == 101 and q.profit_cents == 399


def test_nothing_known_nothing_sold_and_zero_revenue_never_divide_by_zero():
    none_known = profit_summary([_sale(1, [_line("B", 5.5, 2)])])
    assert not none_known.has_profit and none_known.margin is None and none_known.unknown_percent == 100
    assert not none_known.is_partial
    empty = profit_summary([])
    assert empty == ProfitSummary() and empty.margin is None and empty.unknown_percent == 0
    free = profit_summary([_sale(1, [_line("F", 0.0, 4, 3.0)])])  # given away: a known loss, no margin to quote
    assert free.has_profit and free.profit_cents == -1200 and free.margin is None
    assert margin_text(None) == "n/a" and margin_text(39.04) == "39.0%" and margin_text(-60.0) == "-60.0%"


def test_the_unknown_percent_is_never_rounded_to_a_lie():
    tiny_unknown = profit_summary([_sale(1, [_line("A", 1000.0, 1, 400.0), _line("B", 1.0, 1)])])
    assert tiny_unknown.unknown_percent == 1  # 0.1 % unknown still reads 1, not 0
    tiny_known = profit_summary([_sale(1, [_line("A", 1.0, 1, 0.4), _line("B", 1000.0, 1)])])
    assert tiny_known.unknown_percent == 99  # never "100 %" while something is known


def test_summaries_add_up():
    a, b = profit_summary(SALES[:1]), profit_summary(SALES[1:])
    assert a + b == profit_summary(SALES)


def test_a_loss_has_a_negative_margin():
    p = profit_summary([_sale(1, [_line("L", 5.0, 2, 8.0)])])
    assert (p.profit_cents, p.margin) == (-600, -60.0)


# --- the sales report ---------------------------------------------------------


def _report(sales):
    return build_sales_report(sales, datetime(2026, 9, 1), datetime(2026, 10, 1))


def test_sales_report_totals_state_profit_margin_and_the_unknown_share():
    totals = dict(_report(SALES).sections[0].rows)
    assert totals["Revenue"] == "61.00"
    assert totals["Gross profit"] == "19.50" and totals["Gross profit"].number == 19.5
    assert totals["Margin"] == "39.0%" and abs(totals["Margin"].number - 0.39) < 1e-9
    assert totals["Profit note"] == "Cost unknown for 18% of revenue"


def test_sales_report_without_partial_data_has_no_note_and_no_data_says_na():
    known_only = dict(_report([SALES[1]]).sections[0].rows)
    assert known_only["Gross profit"] == "7.50" and known_only["Margin"] == "37.5%" and "Profit note" not in known_only
    unknown_only = dict(_report([_sale(1, [_line("B", 5.5, 2)])]).sections[0].rows)
    assert unknown_only["Gross profit"] == "n/a" and unknown_only["Margin"] == "n/a"
    assert unknown_only["Profit note"] == "Cost unknown for 100% of revenue"
    assert dict(_report([]).sections[0].rows) == {"Sales": "0", "Revenue": "0.00"}  # nothing sold: nothing to say


def test_per_product_breakdown_gains_profit_columns_only_when_something_is_known():
    section = _report(SALES).sections[1]
    assert section.title.startswith("Per-Product Breakdown") and "gross profit" in section.title
    rows = {r[0]: r for r in section.rows}
    assert tuple(rows["A"]) == ("A", "Name A", "3", "30.00", "12.00", "40.0%")
    assert tuple(rows["B"]) == ("B", "Name B", "2", "11.00", "n/a", "n/a")
    assert tuple(rows["C"]) == ("C", "Name C", "1", "20.00", "7.50", "37.5%")

    plain = _report([_sale(1, [_line("B", 5.5, 2)])]).sections[1]
    assert plain.title == "Per-Product Breakdown" and [len(r) for r in plain.rows] == [4]


# --- the period report / the Reports page view ----------------------------------


def test_period_report_totals_and_dealership_ranking_carry_profit():
    period = analytics.period_for("month", date(2026, 9, 3))
    report = analytics.build_period_report(SALES, period, DEALERS)
    totals = dict(report.sections[0].rows)
    assert (totals["Gross profit"], totals["Margin"]) == ("19.50", "39.0%")
    assert totals["Profit note"] == "Cost unknown for 18% of revenue"
    top = next(s for s in report.sections if s.title.startswith("Top dealerships"))
    assert top.title.endswith("gross profit, margin %)")
    by_name = {r[1]: r for r in top.rows}
    assert by_name["Metro Heavy"][3] == "41.00" and by_name["Metro Heavy"][5:] == ("12.00", "40.0%")
    assert by_name["Harbor Point"][5:] == ("7.50", "37.5%")
    assert report.sections[-1].title.startswith("Per-Product Breakdown")


def test_period_report_without_costs_is_unchanged_in_shape():
    period = analytics.period_for("month", date(2026, 9, 3))
    sales = [_sale(1, [_line("B", 5.5, 2)])]
    report = analytics.build_period_report(sales, period, DEALERS)
    top = next(s for s in report.sections if s.title.startswith("Top dealerships"))
    assert top.title == "Top dealerships (rank, name, region, revenue, 7-day trend)" and all(len(r) == 5 for r in top.rows)


def test_report_view_and_top_dealerships_expose_profit():
    period = analytics.period_for("month", date(2026, 9, 3))
    view = analytics.build_report_view(period, SALES, [], DEALERS)
    assert view.profit.profit_cents == 1950 and view.profit.margin == 39.0 and view.profit.unknown_percent == 18
    metro = next(d for d in view.top if d.code == "MET-01")
    assert metro.profit.profit_cents == 1200 and metro.profit.margin == 40.0
    assert analytics.build_report_view(period, [], [], DEALERS).profit == ProfitSummary()


def test_profit_between_uses_the_local_calendar_day_like_revenue_does():
    p = analytics.profit_between(SALES, date(2026, 9, 2), date(2026, 9, 2))
    assert p.profit_cents == 750
    assert analytics.profit_between(SALES, date(2026, 8, 1), date(2026, 8, 31)) == ProfitSummary()


# --- exports ---------------------------------------------------------------------


def _doc():
    period = analytics.period_for("month", date(2026, 9, 3))
    return analytics.build_period_report(SALES, period, DEALERS)


def test_csv_export_includes_profit_rows():
    from admin_app.export.csv_report_exporter import export_to_csv

    with mock.patch("admin_app.export.csv_report_exporter.store_name", lambda: "Shop"):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.csv"
            export_to_csv(_doc(), path)
            rows = list(csv.reader(path.read_text(encoding="utf-8-sig").splitlines()))
    assert ["Gross profit", "19.50"] in rows and ["Margin", "39.0%"] in rows
    assert ["Profit note", "Cost unknown for 18% of revenue"] in rows
    assert ["A", "Name A", "3", "30.00", "12.00", "40.0%"] in rows


def test_excel_export_writes_profit_and_margin_as_real_numbers():
    from openpyxl import load_workbook

    from admin_app.export.excel_report_exporter import export_to_excel

    with mock.patch("admin_app.export.excel_report_exporter.store_name", lambda: "Shop"):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.xlsx"
            export_to_excel(_doc(), path)
            sheet = load_workbook(path).active
    cells = {}
    for row in sheet.iter_rows():
        if row[0].value is not None:
            cells.setdefault(row[0].value, row)
    profit, margin = cells["Gross profit"][1], cells["Margin"][1]
    assert profit.value == 19.5 and profit.number_format == "#,##0.00"
    assert abs(margin.value - 0.39) < 1e-9 and margin.number_format == "0.0%"
    assert cells["Profit note"][1].value == "Cost unknown for 18% of revenue"
    product = cells["A"]
    assert product[4].value == 12.0 and abs(product[5].value - 0.4) < 1e-9
    assert cells["B"][4].value == "n/a" and cells["B"][5].value == "n/a"


def test_pdf_export_carries_the_profit_lines():
    import shutil
    import subprocess

    from admin_app.export.pdf_report_exporter import export_to_pdf

    with mock.patch("admin_app.export.pdf_report_exporter.store_name", lambda: "Shop"):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.pdf"
            export_to_pdf(_doc(), path)
            assert path.read_bytes().startswith(b"%PDF")
            if shutil.which("pdftotext"):
                text = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True).stdout
                assert "Gross profit" in text and "19.50" in text and "Cost unknown for 18% of revenue" in text

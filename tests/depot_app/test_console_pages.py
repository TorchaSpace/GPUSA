"""Offscreen tests: Console Dashboard, Inventory and Reports, and the
Manager Portal showing who is signed in."""

from __future__ import annotations

from datetime import date, datetime, timedelta

import pytest

from tests.gui_support import pump, qapp  # noqa: F401

import database.connection as connection
from database import (
    account_repository,
    attendance_repository,
    dealership_repository,
    employee_repository,
    product_repository,
    purchase_order_repository,
    shipment_repository,
    stock_repository,
    warehouse_repository,
)
from shared import auth, current_session
from shared.models import UNASSIGNED, Dealership, Employee, Product, Warehouse


@pytest.fixture(autouse=True)
def _isolated_db(tmp_path, monkeypatch):
    monkeypatch.setattr(connection, "get_db_path", lambda: tmp_path / "t.db")
    monkeypatch.setattr(connection, "_initialized", False)


@pytest.fixture
def world():
    w = warehouse_repository.create(Warehouse(code="WH-01", name="Merkez", capacity_units=500))
    dealership_repository.create(Dealership(code="001", name="Harbor Point", region="Coastal", city="Norfolk"))
    product_repository.create(Product("BOX", "Carton", 40, 300, 40))
    product_repository.create(Product("TAPE", "Tape", 5, 0, 5))
    # TAPE was stocked here before and is now at zero (a level row exists): THAT is what makes this
    # warehouse alert on it - a product it never carried is not "low" here.
    with connection.connection_scope() as conn:
        conn.execute("INSERT INTO stock_levels (location_kind, location_code, product_barcode, quantity) "
                     "VALUES ('warehouse', 'WH-01', 'TAPE', 0)")
    account_repository.create_first_admin("A-1", "Erol", "482913")
    employee_repository.create(Employee("M-1", "Murat", "Operations", "Warehouse", "merkez"))
    account_repository.create_account("M-1", "depot_manager", "7351")
    session = account_repository.authenticate("M-1", "7351", auth.AREA_DEPOT_CONSOLE, "Depot WH-01")
    current_session.set(session)
    attendance_repository.check_in("M-1")
    stock_repository.transfer(UNASSIGNED, w.location, "BOX", 280, actor=session.actor)
    stock_repository.dispatch(w.location, "BOX", 3, "damaged")
    s = shipment_repository.create(w.site_label, "001", "X", datetime.now() + timedelta(hours=2), [("BOX", 20)],
                                   origin_code="WH-01")
    shipment_repository.dispatch(s.id, actor=session.actor)
    purchase_order_repository.submit("TAPE", "Acme", 10, 1.0, w.site_label)
    return w, session


def test_dashboard_numbers(qapp, world):
    from depot_app.gui.dashboard_page import DashboardPage

    w, _ = world
    page = DashboardPage(w)
    page.reload()
    cells = {key: cell.value_label.text() for key, cell in page.cells.items()}
    assert cells["capacity"] == "51%"  # 257 / 500
    assert cells["skus"] == "1" and cells["low"] == "1"  # tape: 0 here
    assert cells["today"] == "+280 / −23"
    assert cells["shipments"] == "1" and cells["staff"] == "1 / 1" and cells["orders"] == "1"
    assert page.low_table.item(0, 0).text() == "TAPE"
    assert page.recent_table.item(0, 4).text() == "Shipment"


def test_inventory_filters_and_count(qapp, world):
    from depot_app.gui.inventory_page import InventoryPage

    w, _ = world
    page = InventoryPage(w)
    page.reload()
    assert [page.table.item(r, 0).text() for r in range(page.table.rowCount())] == ["BOX", "TAPE"]
    assert page.table.item(0, 5).text() == "40"  # 20 still unassigned + 20 on the truck
    page.set_filter("Out")
    assert [page.table.item(r, 0).text() for r in range(page.table.rowCount())] == ["TAPE"]
    page.set_filter("All")
    page.search_input.setText("cart")
    assert page.table.rowCount() == 1
    assert page.record_count("BOX", 250, "shelf check") == -7
    counted = stock_repository.list_movements(1, location=w.location)[0]
    assert (counted["reason"], counted["handled_by"]) == ("count", "Murat · M-1")
    assert "(-7)" in page.message.text()


def test_reports_and_exports(qapp, world, tmp_path):
    from depot_app.gui.reports_page import ReportsPage

    w, _ = world
    page = ReportsPage(w)
    page.set_last_days(1)
    report = page.report
    assert (report.units_in, report.units_out, len(report.movements)) == (280, 23, 3)
    assert page.cells["net"].value_label.text() == "+257"
    headers = [page.summary_table.horizontalHeaderItem(c).text() for c in range(page.summary_table.columnCount())]
    assert headers == ["SKU", "Product", "In", "Out", "Net", "Written out", "Shipped out", "Transfers in"]
    assert page.detail_table.item(0, 7).text() == "Murat · M-1"
    assert page.excel_button.isEnabled()
    xlsx = page.export_to(tmp_path / "r.xlsx", "xlsx")
    pdf = page.export_to(tmp_path / "r.pdf", "pdf")
    from openpyxl import load_workbook

    book = load_workbook(xlsx)
    assert book.sheetnames == ["Summary", "Movements"] and book["Movements"].max_row == 4
    assert pdf.read_bytes().startswith(b"%PDF")


def test_reports_period_excludes_other_days(qapp, world):
    from depot_app.gui.reports_page import ReportsPage

    w, _ = world
    page = ReportsPage(w, today_provider=lambda: date.today() - timedelta(days=3))
    page.set_last_days(1)
    assert page.report.movements == [] and not page.excel_button.isEnabled()
    page.to_input.setDate(page.from_input.date().addDays(-1))
    assert page.run() is None and "before" in page.message.text()


def test_portal_names_the_signed_in_manager(qapp, world):
    from depot_app.gui.manager_portal_dialog import ManagerPortalDialog

    w, session = world
    portal = ManagerPortalDialog(w.site_label)
    portal.show()
    pump(qapp)
    assert portal._subtext_label.text().startswith("Murat · Depot manager · unlocked")
    assert portal._kicker_label.text() == "İDARİ GİRİŞ · DEPOT MANAGER"
    portal.close()


def test_portal_unlock_checks_the_pin(qapp, world):
    from depot_app.gui.auth_flow import portal_unlock_dialog

    _, session = world
    dialog = portal_unlock_dialog(session)
    assert dialog.badge_input.isReadOnly() and dialog.badge_input.text() == "M-1"
    dialog.pin_input.setText("0000")
    dialog.try_sign_in()
    assert dialog.error_label.text().startswith("Badge or PIN is wrong.")
    dialog.pin_input.setText("7351")
    dialog.try_sign_in()
    assert dialog.session == session
    assert account_repository.list_events(1)[0]["event"] == "pin_confirmed"


def test_inventory_does_not_call_a_never_stocked_product_out(qapp, world):
    from depot_app.gui.inventory_page import InventoryPage, status_here

    w, _ = world
    product_repository.create(Product("GLUE", "Glue", 3, 0, 5))  # this warehouse never carried it
    page = InventoryPage(w)
    page.reload()
    rows = {page.table.item(r, 0).text(): page.table.item(r, 4).text() for r in range(page.table.rowCount())}
    assert rows["GLUE"] == "Not stocked" and rows["TAPE"] == "Out"
    page.set_filter("Out")
    assert [page.table.item(r, 0).text() for r in range(page.table.rowCount())] == ["TAPE"]
    assert status_here(Product("X", "x", 1, 0, 5, stocked_here=False)) == "Not stocked"
    assert status_here(Product("X", "x", 1, 0, 5)) == "Out"
